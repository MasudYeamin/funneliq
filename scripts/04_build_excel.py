"""
FunnelIQ - Excel workbook builder
Produces FunnelIQ_Report.xlsx with a raw-data sheet, a pivot-style manual
funnel summary built with native Excel formulas (SUMPRODUCT/COUNTIFS), and
the A/B test readout - demonstrating the "quick Excel pass before scripting"
workflow.
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference

orders = pd.read_csv("data/raw/orders.csv")
events = pd.read_csv("data/raw/events.csv")
ab_test = pd.read_csv("data/processed/ab_test_results.csv")

OUT_PATH = "excel/FunnelIQ_Report.xlsx"

with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
    events.head(5000).to_excel(writer, sheet_name="raw_events", index=False)
    orders.to_excel(writer, sheet_name="raw_orders", index=False)
    ab_test.to_excel(writer, sheet_name="ab_test_results", index=False)

wb = load_workbook(OUT_PATH)

# ---- Funnel Summary sheet with live Excel formulas (COUNTIFS/SUMPRODUCT) --
ws = wb.create_sheet("Funnel_Summary", 0)
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4472C4")

ws.append(["funnel_step", "customers_reached (COUNTIFS-driven)"])
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill

last_event_row = min(events.shape[0], 5000) + 1  # matches raw_events sheet range

steps = ["view", "add_to_cart", "checkout", "purchase"]
for i, step in enumerate(steps, start=2):
    ws.cell(row=i, column=1, value=step)
    # SUMPRODUCT of unique customer_id count per funnel_step on the raw_events sheet
    formula = (
        f'=SUMPRODUCT((raw_events!$D$2:$D${last_event_row}="{step}")/'
        f'COUNTIFS(raw_events!$D$2:$D${last_event_row},raw_events!$D$2:$D${last_event_row},'
        f'raw_events!$B$2:$B${last_event_row},raw_events!$B$2:$B${last_event_row}))'
    )
    ws.cell(row=i, column=2, value=formula)

chart = BarChart()
chart.title = "Funnel by Step"
chart.y_axis.title = "Customers reached"
data = Reference(ws, min_col=2, min_row=1, max_row=len(steps) + 1)
cats = Reference(ws, min_col=1, min_row=2, max_row=len(steps) + 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "E2")

for col, width in zip("AB", [18, 34]):
    ws.column_dimensions[col].width = width

wb.save(OUT_PATH)
print(f"Saved {OUT_PATH}")
